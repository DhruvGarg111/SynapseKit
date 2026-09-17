from __future__ import annotations

import asyncio
import os

from .base import Document


class PowerPointLoader:
    """Load a PowerPoint (.pptx) file, one Document per slide.

    Extracts text from all shapes (text frames) on each slide.
    """

    def __init__(self, path: str) -> None:
        self._path = path

    async def aload(self) -> list[Document]:
        """Load slide text without blocking the event loop."""
        return await asyncio.to_thread(self.load)

    def load(self) -> list[Document]:
        if not os.path.exists(self._path):
            raise FileNotFoundError(f"PowerPoint file not found: {self._path}")
        try:
            from pptx import Presentation
        except ImportError:
            raise ImportError("python-pptx required: pip install synapsekit[pptx]") from None

        prs = Presentation(self._path)
        docs = []
        for i, slide in enumerate(prs.slides):
            parts = []
            for shape in slide.shapes:
                if getattr(shape, "has_table", False):
                    for row in shape.table.rows:
                        cells = [cell.text.strip() for cell in row.cells]
                        if any(cells):
                            parts.append(" | ".join(cells))
                if getattr(shape, "has_text_frame", False):
                    for paragraph in shape.text_frame.paragraphs:
                        text = paragraph.text.strip()
                        if text:
                            parts.append(text)
            docs.append(
                Document(
                    text="\n".join(parts),
                    metadata={
                        "source": self._path,
                        "file": self._path,
                        "source_type": "pptx",
                        "chunk_type": "slide",
                        "page": i + 1,
                        # ``slide`` is a legacy zero-based field; ``page`` is
                        # the human-facing one-based locator.
                        "slide": i,
                        "slide_number": i + 1,
                        "locator": f"{os.path.basename(self._path)} page {i + 1}",
                    },
                )
            )
        return docs
